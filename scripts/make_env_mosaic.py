"""Main figure, random-layout version with more maps: a gap-free mosaic of N_PER open_cleanup maps (top rows)
and N_PER open_harvest maps (bottom rows), every map a different layout from its env's held-out pool, with
N_AGENTS agents in each. Same drawing as make_env_mosaic_random.py (the 2026-09-18, 3-agent, 5+5 figure), which is
left untouched:
  cleanup  -- one waste hue per map (the evaluated setting, skew 1), hues cycling red/yellow/green/cyan/purple, all
              three shades mixed; a random share of the river already cleared, apples where it is clean enough.
  harvest  -- the baseline since 2026-09-24 (layouts_orig6.py): the original Commons Harvest map, 16 x 22 with six
              patches, each patch's colour drawn at random with replacement, a random share of trees eaten, a random
              mix of ripeness stages.
Every panel is the environment's own renderer; all cells share a width.
Writes docs/figures/env_mosaic_random_{N}ag.{png,pdf}, the single panels (plus crisp 4x copies for slides) in
docs/figures/panels_random_{N}ag/, and docs/env_mosaic_random_{N}ag_<date>.pptx.
usage: python scripts/make_env_mosaic.py [N_AGENTS=5] [N_PER=12] [COLS=6]"""
import os, sys, time
sys.path.insert(0, ".")
import jax, jax.numpy as jnp, numpy as onp
from PIL import Image

N_AGENTS = int(sys.argv[1]) if len(sys.argv) > 1 else 5
N_PER = int(sys.argv[2]) if len(sys.argv) > 2 else 12
COLS = int(sys.argv[3]) if len(sys.argv) > 3 else 6
TAG = f"{N_AGENTS}ag"
OUT = "docs/figures"; PANELS = os.path.join(OUT, f"panels_random_{TAG}"); os.makedirs(PANELS, exist_ok=True)
CELL_W = 765
rng = onp.random.default_rng(31)


def to_cell(img):
    im = Image.fromarray(onp.asarray(img).astype(onp.uint8))
    im = im.resize((im.width * 2, im.height * 2), Image.NEAREST)
    return im.resize((CELL_W, round(CELL_W * im.height / im.width)), Image.LANCZOS)


def save_panel(img, name):
    raw = Image.fromarray(onp.asarray(img).astype(onp.uint8))
    raw.resize((raw.width * 4, raw.height * 4), Image.NEAREST).save(os.path.join(PANELS, f"{name}_big.png"))
    cell = to_cell(img); cell.save(os.path.join(PANELS, f"{name}.png"))
    return cell


def place_agents(state, cells, base_code, floor):
    """Drop the agents on the given free cells (grid + agent_locs), facing at random."""
    g = onp.array(state.grid); locs = onp.array(state.agent_locs)
    for i in range(len(cells)):
        g[locs[i][0], locs[i][1]] = floor
    for i, (r, c) in enumerate(cells):
        g[r, c] = base_code + i; locs[i] = (r, c, int(rng.integers(4)))
    return state.replace(grid=jnp.array(g, dtype=state.grid.dtype), agent_locs=jnp.array(locs, dtype=state.agent_locs.dtype))


def free_cells(g, floor, k, avoid, apart=3):
    """k random cells that are plain ground, not in `avoid`, and (if possible) at least `apart` apart."""
    opts = [(r, c) for r in range(g.shape[0]) for c in range(g.shape[1]) if g[r, c] == floor and (r, c) not in avoid]
    for gap in (apart, apart - 1, 1):
        for _ in range(400):
            pick = [opts[int(i)] for i in rng.choice(len(opts), size=k, replace=False)]
            if all(max(abs(a[0] - b[0]), abs(a[1] - b[1])) >= gap for i, a in enumerate(pick) for b in pick[i + 1:]):
                return pick
    return pick


# ---------------------------------------------------------------- open_cleanup, random layouts
from llm_policy.open_cleanup_policy import make_env as cleanup_env
from opensocialjax.environments.open_cleanup.open_cleanup import Items as CItems, NUM_SHADES, ATTR_HUE_NAMES

cenv = cleanup_env(N_AGENTS, 200, 1, 0.1, 1.0, layouts="random")
DIRTY_PANELS = {COLS // 2 - 2, COLS + COLS // 2 + 1}
top, used, seed = [], set(), 300
for p in range(N_PER):
    while True:                                           # a different map in every panel
        _, st = cenv.reset(jax.random.PRNGKey(seed)); seed += 1
        lid = int(onp.asarray(st.layout_id))
        if lid not in used:
            used.add(lid); break
    lay = cenv.layout_arrays(st); g = onp.array(st.grid); hue = p % 5
    waste = [tuple(int(x) for x in rc) for rc in lay["waste_cells"]]
    # share of the waste already cleared; apples grow below 40 % dirt. Most maps are shown clean enough for apples,
    # two (one per row) still too dirty, so the figure shows both states.
    clean = float(rng.uniform(0.3, 0.45)) if p in DIRTY_PANELS else float(rng.uniform(0.66, 0.92))
    for rc in waste:
        g[rc] = int(CItems.potential_dirt) if rng.random() < clean else 8 + hue * NUM_SHADES + int(rng.choice(3, p=[0.3, 0.35, 0.35]))
    dirt_frac = sum(1 for rc in waste if g[rc] >= 8) / len(waste)
    orchard = sorted(lay["orchard_set"])
    p_apple = max(0.0, (0.4 - dirt_frac) / 0.4) * 0.55
    for rc in orchard:
        if rng.random() < p_apple:
            g[rc] = int(CItems.apple)
    st = st.replace(grid=jnp.array(g, dtype=st.grid.dtype))
    floor = int(CItems.empty)
    st = place_agents(st, free_cells(onp.array(st.grid), floor, N_AGENTS, set(orchard) | set(waste)), len(CItems), floor)
    top.append(save_panel(cenv.render(st), f"cleanup_layout{p}"))
    print(f"cleanup panel {p}: layout {lid}, {ATTR_HUE_NAMES[hue]} waste, {dirt_frac:.0%} of the river dirty, "
          f"{sum(1 for rc in orchard if g[rc] == int(CItems.apple))} apples", flush=True)

# ---------------------------------------------------------------- open_harvest, random layouts
import opensocialjax
from opensocialjax.environments.open_harvest.layouts_orig6 import layout_split
from opensocialjax.environments.open_harvest.open_harvest import Items as HItems, APPLE_BASE, NUM_RIPE_STAGES, HUE_NAMES

pool = layout_split(on_train=False)                        # the held-out maps, the ones an evaluated policy meets
picks = [int(i) for i in rng.choice(pool.size, size=N_PER, replace=False)]
bottom = []
for p, mi in enumerate(picks):
    amap = pool.map_ascii(mi)
    hues6 = tuple(int(h) for h in rng.integers(0, 5, size=6))                 # each of the six patches' colour, drawn with replacement
    henv = opensocialjax.make("open_harvest", num_agents=N_AGENTS, num_inner_steps=500, num_outer_steps=1,
                              fixed_patch_hues=hues6, map_ASCII=amap, grid_size=(len(amap), max(len(r) for r in amap)))
    _, st = henv.reset(jax.random.PRNGKey(400 + p))
    g = onp.array(st.grid); hues = onp.array(st.cell_hues)
    trees = [tuple(int(x) for x in rc) for rc in onp.array(henv.SPAWNS_APPLE)]
    floor = int(HItems.empty)
    eaten = float(rng.uniform(0.05, 0.5)); stages = rng.dirichlet([2, 2, 1.5])
    for k, rc in enumerate(trees):
        g[rc] = floor if rng.random() < eaten else APPLE_BASE + int(hues[k]) * NUM_RIPE_STAGES + int(rng.choice(3, p=stages))
    st = st.replace(grid=jnp.array(g, dtype=st.grid.dtype))
    st = place_agents(st, free_cells(onp.array(st.grid), floor, N_AGENTS, set(trees)), len(HItems), floor)
    bottom.append(save_panel(henv.render(st), f"harvest_layout{p}"))
    print(f"harvest panel {p}: pool map {mi}, patches {'_'.join(HUE_NAMES[h] for h in hues6)}, "
          f"{sum(1 for rc in trees if g[rc] == floor)} of {len(trees)} trees empty", flush=True)

# ---------------------------------------------------------------- stitch: cleanup rows, then harvest rows
rows_c, rows_h = -(-N_PER // COLS), -(-N_PER // COLS)
hc, hh = top[0].height, bottom[0].height
W = CELL_W * COLS; H = rows_c * hc + rows_h * hh
mosaic = Image.new("RGB", (W, H), (0, 0, 0))
for k, im in enumerate(top): mosaic.paste(im, ((k % COLS) * CELL_W, (k // COLS) * hc))
for k, im in enumerate(bottom): mosaic.paste(im, ((k % COLS) * CELL_W, rows_c * hc + (k // COLS) * hh))
png = os.path.join(OUT, f"env_mosaic_random_{TAG}.png"); mosaic.save(png, optimize=True)
mosaic.save(os.path.join(OUT, f"env_mosaic_random_{TAG}.pdf"), resolution=W / 7.0)
print("mosaic", mosaic.size, "->", png)

from pptx import Presentation
from pptx.util import Inches
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
height = Inches(7.1); width = int(height * W / H)
if width > Inches(12.9):
    width = Inches(12.9); height = int(width * H / W)
slide.shapes.add_picture(png, int((prs.slide_width - width) / 2), int((prs.slide_height - height) / 2), width=width)
out_pptx = f"docs/env_mosaic_random_{TAG}_{time.strftime('%Y-%m-%d')}.pptx"; prs.save(out_pptx); print("pptx ->", out_pptx)
