"""A tiny drawing layer for paper figures: every primitive is written twice, as a native PowerPoint shape (editable
.pptx) and onto a PIL image (PNG / PDF preview) from the same geometry in inches. Figures are drawn at print size
(e.g. 7.0 in wide for a full-width figure), so font sizes are the sizes on the page.
Same primitives as scripts/make_llm_agent_figure.py (which keeps its own copy), plus pill() and smooth pictures."""
import math, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

SANS, MONO = "Calibri", "Courier New"
BLACK, GREY, LINE, DIM = (0x22, 0x22, 0x22), (0x59, 0x59, 0x59), (0xA6, 0xA6, 0xA6), (0x9A, 0x9A, 0x9A)
CARD, HEAD, ORANGE, BLUE = (0xF7, 0xF7, 0xF7), (0xE3, 0xEA, 0xF4), (0xD9, 0x6B, 0x00), (0x1F, 0x5F, 0xBF)
PILL, PILL_LINE, PILL_TEXT = (0xFF, 0xF3, 0xD0), (0xBF, 0x90, 0x00), (0x7A, 0x5A, 0x00)
HUE_RGB = {"red": (0xD6, 0x45, 0x45), "yellow": (0xE0, 0xB8, 0x2E), "green": (0x3F, 0xA6, 0x4B), "cyan": (0x2F, 0xB3, 0xC4), "purple": (0x8E, 0x5B, 0xC9)}
FD = "/usr/share/fonts/truetype"
rgb = lambda c: RGBColor(*c)


class Fig:
    def __init__(self, sw, sh, dpi=300):
        self.sw, self.sh, self.dpi = sw, sh, dpi
        self.prs = Presentation(); self.prs.slide_width, self.prs.slide_height = Inches(sw), Inches(sh)
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.im = Image.new("RGB", (int(sw * dpi), int(sh * dpi)), "white"); self.d = ImageDraw.Draw(self.im)

    def font(self, mono, bold, pt):
        name = ("LiberationMono" if mono else "LiberationSans") + ("-Bold" if bold else "-Regular") + ".ttf"
        return ImageFont.truetype(os.path.join(FD, name), max(6, int(pt * self.dpi / 72 * (1.0 if mono else 0.93))))

    def box(self, x, y, w, h, fill=CARD, line=LINE, radius=0.06, width=1.0, dash=False, name=None):
        D = self.dpi
        sh = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius: sh.adjustments[0] = radius
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if line is None: sh.line.fill.background()
        else: sh.line.color.rgb = rgb(line); sh.line.width = Pt(width)
        sh.shadow.inherit = False; sh.text_frame.text = ""
        if dash: sh.line.dash_style = MSO_LINE.DASH
        if name: sh.name = name
        r = int(min(w, h) * radius * D) if radius else 0
        self.d.rounded_rectangle([x * D, y * D, (x + w) * D, (y + h) * D], r, fill=fill, outline=line, width=max(1, int(width * D / 72)) if line else 0)
        return sh

    def text(self, x, y, w, h, runs, size=7.3, mono=False, colour=BLACK, align="l", anchor="t", wrap=True, spacing=1.0):
        """runs: a string, or a list of lines; a line is a string or a list of (text, bold) / (text, bold, colour) runs."""
        D = self.dpi; lines = [runs] if isinstance(runs, str) else runs
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
        tf.word_wrap = wrap; tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = Inches(0.01)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
        lh = size * 1.2 * spacing / 72; yy = y + 0.01 + ((h - lh * len(lines)) / 2 if anchor == "m" else 0)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]; p.line_spacing = spacing
            parts = [(ln, False)] if isinstance(ln, str) else ln
            parts = [(t[0], t[1], t[2] if len(t) > 2 and t[2] is not None else colour, bool(t[3]) if len(t) > 3 else False) for t in parts]   # (text, bold, colour, superscript)
            fs = lambda b, sup: self.font(mono, b, size * (0.68 if sup else 1.0))
            tw = sum(self.d.textlength(t, font=fs(b, sup)) for t, b, _, sup in parts) / D
            xx = x + 0.03 if align == "l" else x + (w - tw) / 2 if align == "c" else x + w - 0.03 - tw
            for t, b, col, sup in parts:
                r = p.add_run(); r.text = t; r.font.name = MONO if mono else SANS; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = rgb(col)
                if sup: r._r.get_or_add_rPr().set("baseline", "30000")
                f = fs(b, sup); self.d.text((xx * D, (yy - (size * 0.10 / 72 if sup else 0)) * D), t, fill=col, font=f); xx += self.d.textlength(t, font=f) / D
            yy += lh
        return tb

    def pill(self, x, y, w, h, label, size=7.0, fill=PILL, line=PILL_LINE, colour=PILL_TEXT, bold=True):
        self.box(x, y, w, h, fill=fill, line=line, radius=0.5, width=0.75)
        self.text(x, y, w, h, [[(label, bold, colour)]], size=size, align="c", anchor="m")

    def picture(self, path, x, y, w, smooth=False):
        D = self.dpi; self.s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w)); im = Image.open(path).convert("RGBA")
        hh = w * im.height / im.width; small = im.resize((int(w * D), int(hh * D)), Image.LANCZOS if smooth else Image.NEAREST)
        self.im.paste(small, (int(x * D), int(y * D)), small)
        return hh

    def line(self, x1, y1, x2, y2, colour=GREY, width=1.25, head=True, dash=False, head_len=0.09):
        D = self.dpi
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(colour); c.line.width = Pt(width)
        if dash: c.line.dash_style = MSO_LINE.DASH
        if head:
            end = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")); end.set("type", "triangle"); end.set("w", "med"); end.set("len", "med")
        px = max(1, int(width * D / 72))
        if dash:
            n = max(1, int(math.hypot(x2 - x1, y2 - y1) / 0.075))
            for k in range(n):
                t0, t1 = k / n, (k + 0.6) / n
                self.d.line([(x1 + (x2 - x1) * t0) * D, (y1 + (y2 - y1) * t0) * D, (x1 + (x2 - x1) * t1) * D, (y1 + (y2 - y1) * t1) * D], fill=colour, width=px)
        else:
            self.d.line([x1 * D, y1 * D, x2 * D, y2 * D], fill=colour, width=px)
        if head:
            ang = math.atan2(y2 - y1, x2 - x1); L = head_len * D
            self.d.polygon([(x2 * D, y2 * D)] + [(x2 * D - L * math.cos(ang + s), y2 * D - L * math.sin(ang + s)) for s in (0.42, -0.42)], fill=colour)
        return c

    def card(self, x, y, w, h, num, title, note="", head_pt=8.2, hh=0.19):
        self.box(x, y, w, h, radius=0.04, width=0.75); self.box(x, y, w, hh, fill=HEAD, line=LINE, radius=0.04, width=0.75)
        self.text(x + 0.03, y, w - 0.06, hh, [[(f"{num}  {title}", True), ("   " + note if note else "", False)]], size=head_pt, anchor="m")
        return y + hh

    def save(self, pptx, png):
        os.makedirs(os.path.dirname(png), exist_ok=True)
        self.im.save(png, optimize=True, dpi=(self.dpi, self.dpi)); self.im.save(png.replace(".png", ".pdf"), resolution=self.dpi); self.prs.save(pptx)
        print("wrote", pptx, "|", png, "|", png.replace(".png", ".pdf"))
